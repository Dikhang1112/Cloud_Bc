import boto3
import pandas as pd
from io import StringIO, BytesIO
import json
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors

def load_s3(access_key, secret_key, bucket, file_path, data):
    try:
        # Khởi tạo S3 client với credentials
        s3_client = boto3.client(
            's3',
            aws_access_key_id=access_key,
            aws_secret_access_key=secret_key,
            region_name='us-east-1'
        )

        # Chuyển data thành DataFrame
        df = pd.DataFrame(data)

        # Xác định định dạng file dựa trên đuôi file
        file_extension = file_path.lower().split('.')[-1]

        if file_extension == 'csv':
            # Lưu dưới dạng CSV
            csv_buffer = StringIO()
            df.to_csv(csv_buffer, index=False)
            body = csv_buffer.getvalue()

        elif file_extension == 'xlsx':
            # Lưu dưới dạng Excel
            excel_buffer = BytesIO()
            df.to_excel(excel_buffer, index=False, engine='openpyxl')
            excel_buffer.seek(0)
            body = excel_buffer.getvalue()

        elif file_extension == 'json':
            # Lưu dưới dạng JSON
            json_buffer = StringIO()
            df.to_json(json_buffer, orient='records', force_ascii=False)
            body = json_buffer.getvalue()

        elif file_extension == 'docx':
            # Lưu dưới dạng DOCX
            doc = docx.Document()
            table = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
            table.style = 'Table Grid'

            # Thêm header
            for i, column in enumerate(df.columns):
                table.cell(0, i).text = str(column)

            # Thêm dữ liệu
            for i, row in df.iterrows():
                for j, value in enumerate(row):
                    table.cell(i + 1, j).text = str(value)

            doc_buffer = BytesIO()
            doc.save(doc_buffer)
            doc_buffer.seek(0)
            body = doc_buffer.getvalue()

        elif file_extension == 'pdf':
            # Lưu dưới dạng PDF
            pdf_buffer = BytesIO()
            doc = SimpleDocTemplate(pdf_buffer, pagesize=letter)
            elements = []

            # Chuyển DataFrame thành dữ liệu dạng bảng cho PDF
            data_for_table = [df.columns.tolist()] + df.values.tolist()
            table = Table(data_for_table)

            # Định dạng bảng
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ]))

            elements.append(table)
            doc.build(elements)
            pdf_buffer.seek(0)
            body = pdf_buffer.getvalue()

        else:
            raise Exception(f"Định dạng file không được hỗ trợ: {file_extension}. Chỉ hỗ trợ: csv, xlsx, json, docx, pdf")

        # Upload file lên S3 với đường dẫn do người dùng chỉ định
        s3_client.put_object(
            Bucket=bucket,
            Key=file_path,
            Body=body
        )

        return True

    except Exception as e:
        raise Exception(f"Lỗi khi lưu file: {str(e)}")